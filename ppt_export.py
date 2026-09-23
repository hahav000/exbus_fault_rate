"""`report_builder.Slide` 리스트를 실제 .pptx 파일(바이트)로 굳힌다.

Plotly 그림 -> PNG는 kaleido로 뽑는다(정적 이미지라 PPT에 들어갈 수 있는 유일한 형태 -
PPT 안에 인터랙티브 차트를 넣을 방법은 없다). 이미지는 Pillow로 실제 픽셀 크기를 읽어
비율을 유지한 채(object-fit: contain과 같은 방식) 배치한다 - 폭/높이를 둘 다 강제로
지정하면 차트가 찌그러지기 때문이다.
"""

from __future__ import annotations

import io

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from report_builder import Slide

# Streamlit Cloud 같은 배포 환경에는 Chrome이 따로 깔려있지 않아, kaleido(1.0+)가 PNG를
# 못 뽑고 ChromeNotFoundError를 낸다(로컬은 이미 Chrome이 있어 안 겪는 문제). kaleido가
# 자체 제공하는 `get_chrome_sync()`로 Chrome for Testing을 한 번만 내려받아 두면 이후
# to_image() 호출이 그 Chrome을 찾아 쓴다 - 이미 받아져 있으면(버전 태그 일치) 즉시
# 반환되므로 매번 불러도 비용이 거의 없다.
_chrome_ready = False


def _ensure_chrome() -> None:
    global _chrome_ready
    if _chrome_ready:
        return
    try:
        import kaleido
        if hasattr(kaleido, "get_chrome_sync"):
            kaleido.get_chrome_sync()
    except Exception:
        pass  # 실패해도 이후 to_image() 호출에서 원래 에러가 그대로 드러난다
    _chrome_ready = True


FONT = "맑은 고딕"
INK = RGBColor(0x0B, 0x0B, 0x0B)
MUTED = RGBColor(0x52, 0x51, 0x4E)
ACCENT = RGBColor(0x2A, 0x78, 0xD6)
RULE = RGBColor(0xE1, 0xE0, 0xD9)

SLIDE_W = Inches(13.333)   # 16:9
SLIDE_H = Inches(7.5)
MARGIN = Inches(0.5)

# 차트를 몇 인치 폭으로 내보낼지 정하는 기준(인치당 픽셀). 전체 폭 슬라이드(약 12.3in)
# 기준으로 렌더 폭이 1600px쯤 되게 잡은 값이다.
# *** 이 값 하나로 렌더 해상도를 정하되, "몇 인치로 찍힐 박스인가"에 비례해서 렌더
# 폭을 정해야 한다 - 이전엔 이미지 개수와 무관하게 항상 1600px로 고정해서, 차트 2개를
# 나란히 넣는 슬라이드(외장 LTE모뎀 전환 등)의 박스는 절반 폭인데 같은 1600px로 찍다
# 보니 그 안의 글자(legend·축)가 실제 인쇄 크기 기준으로 전체 폭 슬라이드의 절반만큼
# 작게 보였다. 박스 폭에 비례해 렌더 폭을 정하면 어느 슬라이드에서든 같은 실제 크기로
# 보인다.
_RENDER_DPI = 1600 / ((SLIDE_W - 2 * MARGIN) / 914400)


def _new_presentation() -> Presentation:
    prs = Presentation()
    # 기본 템플릿은 4:3(10x7.5in) 배치를 담고 있어, 위젯 배치를 바꾸는 레이아웃/플레이스홀더는
    # 쓰지 않고 전부 blank 레이아웃 위에 직접 그린다 - 그래야 16:9로 늘려도 안 어긋난다.
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    return prs


def _blank(prs: Presentation):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _para(tf, text: str, *, size: Pt, bold: bool = False, color: RGBColor = INK,
         align=None, first: bool = True):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.text = text
    p.font.name = FONT
    p.font.size = size
    p.font.bold = bold
    p.font.color.rgb = color
    if align is not None:
        p.alignment = align
    return p


def _add_title_slide(prs: Presentation, title: str, subtitle: str, footer: str) -> None:
    slide = _blank(prs)
    tb = slide.shapes.add_textbox(MARGIN, Inches(2.7), SLIDE_W - 2 * MARGIN, Inches(1.4))
    tb.text_frame.word_wrap = True
    _para(tb.text_frame, title, size=Pt(40), bold=True, align=PP_ALIGN.CENTER)

    sb = slide.shapes.add_textbox(MARGIN, Inches(4.1), SLIDE_W - 2 * MARGIN, Inches(0.7))
    _para(sb.text_frame, subtitle, size=Pt(18), color=MUTED, align=PP_ALIGN.CENTER)

    fb = slide.shapes.add_textbox(MARGIN, SLIDE_H - Inches(0.6), SLIDE_W - 2 * MARGIN,
                                  Inches(0.4))
    _para(fb.text_frame, footer, size=Pt(11), color=MUTED, align=PP_ALIGN.CENTER)


def _fit(img_bytes: bytes, box_w: int, box_h: int) -> tuple[int, int]:
    """이미지를 찌그러뜨리지 않고 box 안에 담을 (width, height)를 계산한다(contain)."""
    im = Image.open(io.BytesIO(img_bytes))
    px_w, px_h = im.size
    aspect = px_w / px_h
    if box_w / box_h > aspect:
        h = box_h
        w = int(h * aspect)
    else:
        w = box_w
        h = int(w / aspect)
    return w, h


def _add_content_slide(prs: Presentation, spec: Slide, index: int, total: int) -> None:
    slide = _blank(prs)

    tb = slide.shapes.add_textbox(MARGIN, Inches(0.22), SLIDE_W - 2 * MARGIN, Inches(0.6))
    _para(tb.text_frame, spec.title, size=Pt(26), bold=True)
    rule = slide.shapes.add_connector(1, MARGIN, Inches(0.85), SLIDE_W - MARGIN, Inches(0.85))
    rule.line.color.rgb = RULE
    rule.line.width = Pt(1)

    n = len(spec.figs)
    widths = spec.fig_widths or [1.0] * n
    total_w = sum(widths)
    gap = Inches(0.3) if n > 1 else Inches(0)
    area_w = int(SLIDE_W - 2 * MARGIN - gap * (n - 1))

    img_top = Inches(1.05)
    img_box_h = int(SLIDE_H - img_top - Inches(1.9))  # 아래쪽에 bullet 영역 확보

    x = int(MARGIN)
    bottoms = []
    for fig, w in zip(spec.figs, widths):
        box_w = int(area_w * (w / total_w))
        # 차트를 박스와 같은 가로세로 비율로 렌더링해야 여백 없이 꽉 찬다 - Plotly 기본
        # width(700px 안팎)로 그대로 뽑으면 정사각형에 가까워 landscape 박스에서 옆에
        # 큰 여백이 남는다. 렌더 폭은 박스의 실제 인치 폭에 비례시킨다(_RENDER_DPI) -
        # 그래야 차트가 1개짜리 슬라이드든 2개가 나란한 슬라이드든 글자가 같은 실제
        # 크기로 보인다.
        render_w = max(200, round(_RENDER_DPI * (box_w / 914400)))
        render_h = max(1, round(render_w / (box_w / img_box_h)))
        png = fig.to_image(format="png", width=render_w, height=render_h, scale=2)
        disp_w, disp_h = _fit(png, box_w, img_box_h)
        left = x + (box_w - disp_w) // 2
        top = int(img_top) + (img_box_h - disp_h) // 2
        slide.shapes.add_picture(io.BytesIO(png), Emu(left), Emu(top),
                                 width=Emu(disp_w), height=Emu(disp_h))
        bottoms.append(top + disp_h)
        x += box_w + int(gap)

    bullet_top = Emu(max(bottoms) + int(Inches(0.15)))
    bb = slide.shapes.add_textbox(MARGIN, bullet_top, SLIDE_W - 2 * MARGIN,
                                  SLIDE_H - bullet_top - Inches(0.3))
    bb.text_frame.word_wrap = True
    for i, b in enumerate(spec.bullets):
        _para(bb.text_frame, f"•  {b}", size=Pt(14), color=MUTED, first=(i == 0))

    pn = slide.shapes.add_textbox(SLIDE_W - Inches(1.2), SLIDE_H - Inches(0.4), Inches(0.9),
                                  Inches(0.3))
    _para(pn.text_frame, f"{index}/{total}", size=Pt(10), color=MUTED, align=PP_ALIGN.RIGHT)


def build_pptx(slides: list[Slide], *, title: str, subtitle: str, footer: str = "") -> bytes:
    """Slide 리스트를 .pptx 바이트로 만든다. 표지 1장 + 내용 슬라이드 순서대로."""
    _ensure_chrome()
    prs = _new_presentation()
    _add_title_slide(prs, title, subtitle, footer)
    for i, s in enumerate(slides, start=1):
        _add_content_slide(prs, s, i, len(slides))
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()
